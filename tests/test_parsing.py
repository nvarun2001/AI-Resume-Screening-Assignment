from docx import Document

from resume_screener.parsing import _find_email, _find_github, parse_resume
from pdf_helper import make_pdf


def test_find_email():
    assert _find_email("Reach me at jane.doe@example.com today") == "jane.doe@example.com"


def test_find_email_none():
    assert _find_email("no address here") is None


def test_find_github_username_only():
    user, url = _find_github("Profile: https://github.com/janedoe")
    assert user == "janedoe"
    assert url == "https://github.com/janedoe"


def test_find_github_strips_repo_path():
    user, url = _find_github("Code at github.com/janedoe/rag-agent/tree/main")
    assert user == "janedoe"
    assert url == "https://github.com/janedoe"


def test_find_github_skips_reserved_path():
    user, _ = _find_github("See github.com/features and github.com/realuser")
    assert user == "realuser"


def test_find_github_none():
    user, url = _find_github("no profile linked")
    assert user is None
    assert url is None


def test_parse_txt_end_to_end(tmp_path):
    path = tmp_path / "candidate.txt"
    path.write_text(
        "Jane Doe\nPython, LangChain\njane@example.com\ngithub.com/janedoe",
        encoding="utf-8",
    )
    result = parse_resume(path)
    assert result.status == "ok"
    assert result.email == "jane@example.com"
    assert result.github_username == "janedoe"
    assert "LangChain" in result.text


def test_parse_docx_end_to_end(tmp_path):
    path = tmp_path / "candidate.docx"
    doc = Document()
    doc.add_paragraph("Asha Rao")
    doc.add_paragraph("Skills: Python, FastAPI")
    doc.add_paragraph("asha@example.com github.com/asharao")
    doc.save(str(path))
    result = parse_resume(path)
    assert result.status == "ok"
    assert result.email == "asha@example.com"
    assert result.github_username == "asharao"
    assert "FastAPI" in result.text


def test_parse_pdf_end_to_end(tmp_path):
    path = tmp_path / "candidate.pdf"
    path.write_bytes(make_pdf("John Smith Python RAG john@example.com github.com/jsmith"))
    result = parse_resume(path)
    assert result.status == "ok"
    assert result.email == "john@example.com"
    assert result.github_username == "jsmith"


def test_corrupt_pdf_does_not_raise(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4 this is not a real pdf body")
    result = parse_resume(path)
    assert result.status == "failed"
    assert result.error is not None


def test_empty_text_marks_failed(tmp_path):
    path = tmp_path / "blank.txt"
    path.write_text("   \n  ", encoding="utf-8")
    result = parse_resume(path)
    assert result.status == "failed"
