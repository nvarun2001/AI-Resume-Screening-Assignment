"""Raw text extraction from resume files plus regex for atomic fields.

Text extraction is format-specific (PDF/DOCX/TXT). Email and GitHub are pulled
with regex rather than the LLM because their shape is fixed and deterministic
extraction is both cheaper and more reliable. Any single-file failure is
returned as a status, never raised, so one bad resume can't stop the batch.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

import pdfplumber
from docx import Document

logger = logging.getLogger(__name__)

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")

# Matches github.com/<user> and captures the first path segment (the username),
# ignoring a trailing repo path, query string, or www prefix.
GITHUB_RE = re.compile(
    r"github\.com/([A-Za-z0-9](?:[A-Za-z0-9\-]{0,37}[A-Za-z0-9])?)",
    re.IGNORECASE,
)

# GitHub reserves these first-segment paths; they are never usernames.
GITHUB_RESERVED = {
    "about",
    "features",
    "pricing",
    "marketplace",
    "explore",
    "topics",
    "sponsors",
    "settings",
    "notifications",
    "orgs",
    "apps",
}


@dataclass
class ParsedResume:
    """Text and atomic fields extracted from a single file."""

    source_file: str
    text: str = ""
    email: str | None = None
    github_url: str | None = None
    github_username: str | None = None
    status: str = "ok"  # ok | failed
    error: str | None = None


def parse_resume(path: str | Path) -> ParsedResume:
    """Extract text and atomic fields from one resume file."""

    path = Path(path)
    result = ParsedResume(source_file=path.name)
    try:
        result.text = _extract_text(path)
    except Exception as exc:  # noqa: BLE001 - one bad file must not stop the batch
        logger.warning("Failed to parse %s: %s", path.name, exc)
        result.status = "failed"
        result.error = f"{type(exc).__name__}: {exc}"
        return result

    if not result.text.strip():
        result.status = "failed"
        result.error = "No extractable text (possibly a scanned/image-only file)"
        return result

    result.email = _find_email(result.text)
    username, url = _find_github(result.text)
    result.github_username = username
    result.github_url = url
    return result


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {suffix}")


def _extract_pdf(path: Path) -> str:
    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    document = Document(str(path))
    parts = [para.text for para in document.paragraphs]
    # Tables often hold skills/contact info in resume templates.
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _find_email(text: str) -> str | None:
    match = EMAIL_RE.search(text)
    return match.group(0) if match else None


def _find_github(text: str) -> tuple[str | None, str | None]:
    for match in GITHUB_RE.finditer(text):
        username = match.group(1)
        if username.lower() in GITHUB_RESERVED:
            continue
        return username, f"https://github.com/{username}"
    return None, None
